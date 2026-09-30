r"""``export/`` 对抗 fuzz——EPUB/DOCX 双语插译写出的结构与协议不变量。

三层 + 定向缺陷回归钉。raw-text 元素内 ``<``/``&``、注释 ``--``/尾 ``-``、
未终结 PI、``<p<div>`` 形错位 tag 等直通面已修并入池；仍避开的是改变
sent/source 文本关系或需额外簿记的面（``<n>`` 哨兵形字面、
``../`` manifest href、fragment 文档、doctype 内部子集、超深嵌套、源侧
``texlate-zh`` 类名）——由文件尾的定向回归测试覆盖。

L1 纯函数（markers/filters/common）：

- ``reconcile_markers``：每个签发 token 在输出中恰出现一次（重复剥除 +
  丢失补尾）；源文自有字面 token（出现于 sent 但未签发）逐处保留；
  既未签发也不在 sent 的 token 形一律剥净；回复本已一致时逐字节原样
  返回；``None`` 回复不炸。
- ``split_on_markers``：片段值拼接恒还原输入；marker 段值 ∈ 给定 token 集。
- ``normalize_text``/``sanitize_xml_text``：幂等；输出无 XML 非法字符、
  无零宽字符、无连续/首尾空白。
- ``safe_language``：任意输入不炸；非 ``None`` 输出恒 ``[\w-]+`` fullmatch。
- ``Ordinals.allocate``：永不在 ``occupied``/``taken`` 里发已占用 token；
  ``marker_name`` 输出恒 ``[A-Z_]+``；``find_markers`` 全部命中
  ``MARKER_RE.fullmatch`` 且按首现序去重。
- ``is_special_text``/``is_apparatus_text``/``is_wordless``/
  ``is_placeholder_only``：任意对抗串不炸、返回 bool。
- ``coerce_glossary``：mapping → ``Glossary``；缺席路径 → ``ExportError``。

L2 单元级（``iter_units``/``insert_translation``/``insert_after`` 直喂）：

- ``iter_units``（epub）对任意 html.parser 可解析的 hostile body 不抛；
  unit.text 已归一化且非空、非 special/apparatus；body 单元非纯占位符；
  job_id 全本唯一；签发 token 在 unit.text 中恰一次、形合 ``MARKER_RE``。
- ``insert_translation`` 对臆造/丢弃/重复 marker 与 markup/控制字符译文
  不抛（``_apply`` 无逐条护栏——一次 raise 炸掉整次导出）；插后文档仍
  well-formed XML；签发 token 不以字面残留在 ``.texlate-zh`` 节点文本；
  NCX 单元逐字节写 ``原文 / sanitize(reconcile(译文))``。
- ``insert_after``（docx）克隆段紧跟源段、子元素只剩 ``w:pPr``/``w:r``、
  剥 ``w14:paraId``/``w14:textId``/``w:sectPr``、译文 run 带
  ``w:color=555555`` + ``w:lang`` 章；对抗译文不炸、part XML 仍合法。

L3 端到端（``translate_epub``/``translate_docx`` 正常返回时）：

- EPUB：``sniff_format`` 回 ``"epub"``；``testzip`` 通过；``mimetype``
  首条 ZIP_STORED 且内容恰为 ``application/epub+zip``；其余成员 DEFLATED；
  成员集 == 输入成员集 ∪ {mimetype}；每个文档成员 + OPF + NCX 均
  well-formed XML；文档开头 xml decl 的 encoding 值恒为 utf-8；
  ``dc:language`` == ``safe_language(target_lang)``（None 则保持原值）；
  NCX ``<text>`` 要么原样要么 ``原规范文本 + " / " + 译文``；
  ``translated`` == 文档侧 ``.texlate-zh`` 节点数 + NCX 被改写数；
  ``translated+unchanged+skipped+fault == units == len(iter_units 重算)``；
  签发 token 不以字面残留译文节点（只有源文自有字面 token 可合法出现）；
  每个 unit 源文的非 marker 片段仍是输出文档可见文本的子串；
  ``on_result`` 每单元恰好回调一次；成功后 ``{dst}.state`` 已清理。
- DOCX：``sniff_format`` 回 ``"docx"``；全部 ``*.xml``/``*.rels`` 成员
  well-formed；源段按原序完整保持为非译文段子序列；译文段
  （``w:color=555555`` 戳）在 ``w:p`` 文档序中紧跟一个非译文段；
  译文段总数 == ``translated``；``documents`` == 1 + 有单元的附属 part
  数；safe ``target_lang`` 时译文 run 带 ``w:lang @w:val=<lang>`` 且
  core language == lang；计数不变量同上。
- 半成品：auth 熔断抛 ``AuthTrippedError`` 后 ``dst`` 仍是合法出包
  （OCF 首条/strict XML/可 sniff），``{dst}.state`` 保留供续跑。
- ``export_document`` 按内容嗅探分派（docx 身 epub 名 → docx 管线）；
  ``sniff_format`` 对任意成员集不炸且 epub 判定蕴含规范 mimetype。
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from contextlib import suppress
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

import pytest
from _exportkit import _zip_only
from _fuzzkit import fuzz_rng
from bs4 import BeautifulSoup
from bs4.element import NavigableString
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.text.paragraph import Paragraph
from lxml import etree

from texlate.export import export_document, sniff_format
from texlate.export.common import (
    ExportError,
    ExportReport,
    UnsupportedFormatError,
    coerce_glossary,
    safe_language,
)
from texlate.export.docx import insert_after, translate_docx
from texlate.export.docx import iter_units as iter_docx_units
from texlate.export.epub import (
    EpubBook,
    Unit,
    book_soups,
    insert_translation,
    iter_units,
    load_epub,
    translate_epub,
)
from texlate.export.filters import (
    is_apparatus_text,
    is_special_text,
    normalize_text,
    sanitize_xml_text,
)
from texlate.export.markers import (
    MARKER_RE,
    Ordinals,
    find_markers,
    is_wordless,
    marker_name,
    marker_token,
    reconcile_markers,
    split_on_markers,
)
from texlate.xlat.client import AuthError
from texlate.xlat.pipeline import AuthTrippedError, ChunkResult, MockTranslator
from texlate.xlat.placeholders import ANY_PH_RX, is_placeholder_only

if TYPE_CHECKING:
    import random
    from collections.abc import Callable, Iterable
    from pathlib import Path

    from lxml.etree import _Element

# ------------------------------------------------------------------ 常量

_ZH_LANG = "zh-CN"
_ITERS_EPUB = 60
_ITERS_DOCX = 40
_ITERS_INSERT = 50
_ITERS_UNITS = 80
_ITERS_PURE = 400
_ITERS_SNIFF = 100
_DEEP_NEST = 2000
_PARAID_EVERY = 5  # 每隔几段打 w14:paraId 锚（克隆必须剥）
_P_SNIFF_MIME = 0.5  # sniff fuzz 里 mimetype 成员内容为规范值的概率

_CONTAINER_PATH = "META-INF/container.xml"
_MIMETYPE = b"application/epub+zip"
_EPUB_NS_DECL = 'xmlns:epub="http://www.idpf.org/2007/ops"'
_DC_LANG_TAG = "{http://purl.org/dc/elements/1.1/}language"
_NUM_LINE_RE = re.compile(r"^(\[\d+\])\s?(.*)$", re.DOTALL)
_DECL_ENCODING_RE = re.compile(rb'encoding\s*=\s*["\']([^"\']+)')
_DOCX_TEXT_PART_RE = re.compile(
    r"word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml"
)
_XML_ILLEGAL_CHARS = re.compile(
    "[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f\\ud800-\\udfff\\ufffe\\uffff]"
)
_INVISIBLE = "­\u200b﻿"

#: 纯散文文本池——可含 marker 形字面（只许落在散文位，不进 code/sup 等可成
#: marker 克隆的元素内，否则克隆文本里的 token 形会误报"签发 token 残留"）。
_PROSE_POOL = [
    "Normal English sentence here.",
    "Another para with enough words to be prose.",
    "短中文句子也需要足够长。",
    "Mixed 中英 prose text with numbers 12345 inside.",
    "Ends with marker literal [[IMG_1]] right here.",
    "Bare token [[SL]] and [[PL]] literals.",
    "Invented-looking [[FAKE_9]] literal in source.",
    "Bracket [[A]] short literal and [[LONG_NAME_22]] too.",
    "Text with ]]&gt; escaped cdata close inside.",
    "Chars &amp; entities &copy; &lt;tag&gt; galore.",
    "RTL نص عربي mixed with עברית tail words.",
    "Combining marks é å ö and CJK 漢字仮名.",
    "Zero­width\u200bjoiner soft­hyphen inside words.",
    "Emoji 🎉🚀 and symbols ∀∂∫∮ math text.",
    "Tab\there and non-breaking space inside.",
    "DEL\x7f char and U+2028 separator.",
    "Figure 3",
    "Listing 12",
    "Source: adapted from somewhere",
    "ISBN 9780132350884",
    "https://example.com/some/very/long/url/path",
    "12345",
    "!!!???...",
    "a b c",
    "x",
    "",
    "   ",
    "\x01\x02control\x0bchars\x00inside",
    "word " * 40,
    "non­breaking hyphenated text runs",
]

#: 元素内容池——进 code/sup/svg-text/ruby-rt 等可成 marker 克隆的元素内部，
#: 禁 ``[[`` 形（克隆内字面 token 无法与"未剥净"区分）。
_ELEM_POOL = [
    "elem text run",
    "短元素文本",
    "x+y=z",
    "foo(bar)",
    "note 7",
    "注音かな",
    "fallback prose here",
    "option one",
    "＊※",
    "短",
]

#: 注释内容池——``--``/尾 ``-`` 由 ``_sanitize_dom`` 折叠净化，可入池；
#: ``<!--`` 仍禁（会提前闭合注释改变输入结构）。
_COMMENT_POOL = [" note text ", "TODO fix later", "中注", "a -- b tail-"]
#: CDATA 内容池——禁 ``]]>``（提前闭合）；``<``/``&`` 在 CDATA 内合法。
_CDATA_POOL = ["raw <b> not markup", "a && b || c", "x < y > z"]

_INLINE_STYLE_POOL = [
    "display:none",
    "color:red",
    "display:block",
    "font-style:italic",
]
_CLASS_POOL = ["c1", "note", "x y", "正文"]
_EPUB_TYPE_POOL = ["pagebreak", "noteref", "footnote", "toc", "page-list"]
_ROLE_POOL = ["doc-pagebreak", "doc-endnote", "doc-tip"]
_TITLE_POOL = ["a title", "x&quot;y", "标题"]
_HREF_POOL = [
    "ch1.xhtml",
    "ch2.xhtml#sec",
    "other doc.xhtml",
    "javascript:void(0)",
    "#frag",
]
_NCX_LABEL_POOL = ["Chapter One", "第二章", "Intro [[SL]] here", "标签 [[X_2]]"]

# ------------------------------------------------------------------ EPUB 构造


def _container_xml(opf_path: str) -> str:
    return f"""<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="{opf_path}" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""


def _opf(items: list[tuple[str, str, str]], spine_refs: list[str]) -> str:
    rows = "\n".join(
        f'    <item id="{iid}" href="{href}" media-type="{mtype}"/>'
        for iid, href, mtype in items
    )
    refs = "\n".join(f'    <itemref idref="{r}"/>' for r in spine_refs)
    return f"""<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="bid">fuzz</dc:identifier>
    <dc:language>en</dc:language>
  </metadata>
  <manifest>
{rows}
  </manifest>
  <spine>
{refs}
  </spine>
</package>
"""


def _ncx(labels: list[str]) -> str:
    navs = "\n".join(
        f'    <navPoint id="np{i}"><navLabel><text>{t}</text></navLabel>'
        f'<content src="ch1.xhtml#s{i}"/></navPoint>'
        for i, t in enumerate(labels)
    )
    return f"""<?xml version="1.0"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <navMap>
{navs}
  </navMap>
</ncx>
"""


def _xhtml(
    body: str,
    *,
    decl: str | None = "utf-8",
    head: str = "",
    doctype: bool = False,
) -> bytes:
    d = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        if decl == "utf-8"
        else '<?xml version="1.0" encoding="ISO-8859-1"?>\n'
        if decl == "iso"
        else '﻿<?xml version="1.0"?>\n'
        if decl == "bom"
        else ""
    )
    dt = "<!DOCTYPE html>" if doctype else ""
    return (
        f'{d}{dt}<html {_EPUB_NS_DECL} xmlns="http://www.w3.org/1999/xhtml">'
        f"<head><title>t</title>{head}</head><body>{body}</body></html>"
    ).encode()


def zip_bytes(members: Iterable[tuple[str, bytes]]) -> bytes:
    """``(成员名，字节)`` 序列写 zip——``dict.items()`` 直喂，重复名亦可表达。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, blob in members:
            z.writestr(name, blob)
    return buf.getvalue()


def _write_epub(tmp: Path, name: str, members: dict[str, bytes]) -> Path:
    p = tmp / f"{name}.epub"
    p.write_bytes(_zip_only(members))
    return p


def _book_soups(book: EpubBook) -> dict[str, BeautifulSoup]:
    """与 ``translate_epub`` 同构的 soup 表（同一 ``book_soups`` 构造口）。"""
    return book_soups(book)


def _single_doc_book(
    ch1: bytes, *, ncx_labels: list[str] | None = None
) -> dict[str, bytes]:
    """单文档 + img 最小书骨架 (可选 ncx)——``_mini_book``/``_pin_book`` 共用。"""
    items = [
        ("c0", "ch1.xhtml", "application/xhtml+xml"),
        ("img", "i.png", "image/png"),
    ]
    members = {
        "mimetype": _MIMETYPE,
        _CONTAINER_PATH: _container_xml("OEBPS/content.opf").encode(),
        "OEBPS/ch1.xhtml": ch1,
        "OEBPS/i.png": b"\x89PNG",
    }
    if ncx_labels is not None:
        items.append(("ncx", "toc.ncx", "application/x-dtbncx+xml"))
        members["OEBPS/toc.ncx"] = _ncx(ncx_labels).encode()
    members["OEBPS/content.opf"] = _opf(items, ["c0"]).encode()
    return members


def _mini_book(body: str, *, ncx_labels: list[str] | None = None) -> dict[str, bytes]:
    """L2 直喂用最小书：单文档 + img，可选 ncx。"""
    return _single_doc_book(_xhtml(body), ncx_labels=ncx_labels)


# ------------------------------------------------------------------ 生成器（EPUB）


def _gen_attrs(rng: random.Random, *, allow_epub_type: bool = True) -> str:
    """随机属性组（值取自安全池——attr 名/值直通面由 ``_sanitize_dom`` 净化，
    本组仍只放合法值以保 oracle 可复算）。每项按独立概率附加。"""
    table: list[tuple[float, Callable[[], str]]] = [
        (0.25, lambda: f'class="{rng.choice(_CLASS_POOL)}"'),
        (0.15, lambda: f'id="e{rng.randrange(10**6)}"'),
        (0.08, lambda: "hidden"),
        (0.12, lambda: f'style="{rng.choice(_INLINE_STYLE_POOL)}"'),
        (0.08, lambda: 'dir="rtl"'),
        (0.08, lambda: f'lang="{rng.choice(["es", "fr", "zh-Hant"])}"'),
        (0.08, lambda: f'title="{rng.choice(_TITLE_POOL)}"'),
        (0.06, lambda: f'role="{rng.choice(_ROLE_POOL)}"'),
        (0.05, lambda: 'aria-hidden="true"'),
    ]
    if allow_epub_type:
        table.append((0.10, lambda: f'epub:type="{rng.choice(_EPUB_TYPE_POOL)}"'))
    attrs = [gen() for prob, gen in table if rng.random() < prob]
    return (" " + " ".join(attrs)) if attrs else ""


_MAX_NEST = 3


def _il_wrap(rng: random.Random, depth: int) -> str:
    tag = rng.choice(
        [
            "span",
            "em",
            "strong",
            "b",
            "i",
            "u",
            "s",
            "mark",
            "small",
            "abbr",
            "cite",
            "q",
            "dfn",
            "time",
            "var",
            "kbd",
            "samp",
            "bdi",
            "bdo",
            "ins",
        ]
    )
    inner = "".join(_gen_inline(rng, depth + 1) for _ in range(rng.randint(1, 3)))
    return f"<{tag}{_gen_attrs(rng)}>{inner}</{tag}>"


def _il_exclude(rng: random.Random, _depth: int) -> str:
    tag = rng.choice(["code", "sup", "sub"])  # exclude 系 → marker 候选
    return f"<{tag}{_gen_attrs(rng)}>{rng.choice(_ELEM_POOL)}</{tag}>"


def _il_link(rng: random.Random, depth: int) -> str:
    inner = "".join(_gen_inline(rng, depth + 1) for _ in range(rng.randint(1, 2)))
    return f'<a href="{rng.choice(_HREF_POOL)}"{_gen_attrs(rng)}>{inner}</a>'


def _il_img(rng: random.Random, _depth: int) -> str:
    src = rng.choice(["i.png", "pic.png", "missing.png"])
    alt = rng.choice(["", "a pic", "图"])
    return f'<img src="{src}" alt="{alt}"{_gen_attrs(rng)}/>'


def _il_ruby(rng: random.Random, _depth: int) -> str:
    base = rng.choice(_ELEM_POOL)
    note = rng.choice(_ELEM_POOL)
    rp = rng.choice(["", "<rp>(</rp>", "<rp>（</rp>"])
    return f"<ruby>{base}<rt>{note}</rt>{rp}</ruby>"


def _il_protected(rng: random.Random, _depth: int) -> str:
    tag = rng.choice(["svg", "math", "iframe", "video", "audio", "canvas", "object"])
    inner = (
        f"<text>{rng.choice(_ELEM_POOL)}</text>"
        if tag == "svg"
        else rng.choice(_ELEM_POOL)
    )
    return f"<{tag}{_gen_attrs(rng)}>{inner}</{tag}>"


#: script/style 变体里"安全内容"占比（另一半是 ``<``/``&`` 对抗形）
_P_RAW_SAFE = 0.5

#: 行内发生器权重表（累计点即概率阈值——表数据免 PLR2004 提名噪音）
_INLINE_GENS: list[tuple[float, Callable[[random.Random, int], str]]] = [
    (0.42, lambda rng, _d: rng.choice(_PROSE_POOL)),
    (0.08, _il_wrap),
    (0.06, _il_exclude),
    (0.06, _il_link),
    (0.06, _il_img),
    (0.04, lambda rng, _d: rng.choice(["<br/>", "<wbr/>", "<hr/>"])),
    (0.04, _il_ruby),
    (0.04, _il_protected),
    (0.04, lambda rng, _d: f"<!--{rng.choice(_COMMENT_POOL)}-->"),
    (0.03, lambda rng, _d: f"<![CDATA[{rng.choice(_CDATA_POOL)}]]>"),
    (0.02, lambda rng, _d: f"<?pi v{rng.randrange(9)} d?>"),
    (0.01, lambda rng, _d: f"<?pi v{rng.randrange(9)} d>"),  # 未终结 PI
    (
        0.03,
        lambda rng, _d: (
            f"<script>var x = {rng.randrange(99)};</script>"
            if rng.random() < _P_RAW_SAFE
            else "<script>if (a < b && c) { go(); }</script>"
        ),
    ),
    (
        0.03,
        lambda rng, _d: (
            f"<style>.a{{color:rgb({rng.randrange(9)},1,1)}}</style>"
            if rng.random() < _P_RAW_SAFE
            else "<style>a > b && c { x: 1 }</style>"
        ),
    ),
    # 结构性错位碎片——html.parser 宽容面（``<p<div>`` 错位 tag 名由
    # ``_sanitize_dom`` 整形）
    (0.03, lambda rng, _d: rng.choice(["</p>", "<b>", "</div>", "<>", "<p<div>"])),
    (0.02, lambda rng, _d: rng.choice(_PROSE_POOL)),
]


def _gen_inline(rng: random.Random, depth: int = 0) -> str:
    """行内片段：散文叶 / 行内元素 / 保护元素 / 渲染空元素。"""
    if depth > _MAX_NEST:
        return rng.choice(_PROSE_POOL)
    cut = rng.random()
    acc = 0.0
    for weight, gen in _INLINE_GENS:
        acc += weight
        if cut < acc:
            return gen(rng, depth)
    return rng.choice(_PROSE_POOL)


def _bk_textblock(rng: random.Random, _depth: int) -> str:
    tag = rng.choice(
        [
            "p",
            "p",
            "p",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "blockquote",
            "address",
            "dt",
            "dd",
        ]
    )
    inner = "".join(_gen_inline(rng) for _ in range(rng.randint(0, 5)))
    return f"<{tag}{_gen_attrs(rng)}>{inner}</{tag}>"


_P_MIXED = 0.4


def _bk_container(rng: random.Random, depth: int) -> str:
    tag = rng.choice(
        [
            "div",
            "section",
            "article",
            "aside",
            "main",
            "header",
            "footer",
            "fieldset",
            "form",
        ]
    )
    inner = "".join(_gen_block(rng, depth + 1) for _ in range(rng.randint(0, 3)))
    if rng.random() < _P_MIXED:  # block+ 裸文本混排 → multi-run owner
        inner = rng.choice(_PROSE_POOL) + inner + rng.choice(_PROSE_POOL)
    return f"<{tag}{_gen_attrs(rng)}>{inner}</{tag}>"


def _bk_list(rng: random.Random, _depth: int) -> str:
    tag = rng.choice(["ul", "ol"])
    items = "".join(
        f"<li{_gen_attrs(rng)}>"
        f"{''.join(_gen_inline(rng) for _ in range(rng.randint(0, 3)))}</li>"
        for _ in range(rng.randint(1, 4))
    )
    return f"<{tag}{_gen_attrs(rng)}>{items}</{tag}>"


def _bk_table(rng: random.Random, _depth: int) -> str:
    rows = "".join(
        "<tr>"
        + "".join(
            f"<td>{''.join(_gen_inline(rng) for _ in range(rng.randint(0, 2)))}</td>"
            for _ in range(rng.randint(1, 3))
        )
        + "</tr>"
        for _ in range(rng.randint(1, 3))
    )
    return f"<table{_gen_attrs(rng)}>{rows}</table>"


def _bk_figure(rng: random.Random, _depth: int) -> str:
    fig = (
        '<img src="i.png" alt="f"/>'
        f"<figcaption>{''.join(_gen_inline(rng) for _ in range(rng.randint(1, 3)))}</figcaption>"
    )
    return f"<figure{_gen_attrs(rng)}>{fig}</figure>"


def _bk_details(rng: random.Random, _depth: int) -> str:
    return (
        f"<details{_gen_attrs(rng)}><summary>{rng.choice(_ELEM_POOL)}</summary>"
        f"<p>{''.join(_gen_inline(rng) for _ in range(rng.randint(1, 3)))}</p></details>"
    )


_P_NAV_LINK = 0.6


def _bk_nav(rng: random.Random, _depth: int) -> str:
    items = "".join(
        f'<li><a href="ch{rng.randint(1, 2)}.xhtml">{rng.choice(_ELEM_POOL)}</a></li>'
        if rng.random() < _P_NAV_LINK
        else f"<li>{rng.choice(_ELEM_POOL)}</li>"
        for _ in range(rng.randint(1, 3))
    )
    return (
        f'<nav epub:type="doc-toc"{_gen_attrs(rng, allow_epub_type=False)}>'
        f"<ol>{items}</ol></nav>"
    )


def _bk_pre(rng: random.Random, _depth: int) -> str:
    return f"<pre{_gen_attrs(rng)}>{rng.choice(_ELEM_POOL)}\n{rng.choice(_ELEM_POOL)}</pre>"


def _bk_dl(rng: random.Random, _depth: int) -> str:
    return (
        f"<dl{_gen_attrs(rng)}><dt>{rng.choice(_ELEM_POOL)}</dt>"
        f"<dd>{''.join(_gen_inline(rng) for _ in range(rng.randint(1, 2)))}</dd></dl>"
    )


def _bk_blockquote(rng: random.Random, _depth: int) -> str:
    inner = "".join(_gen_inline(rng) for _ in range(rng.randint(1, 3)))
    return f"<blockquote{_gen_attrs(rng)}><p>{inner}</p></blockquote>"


def _bk_select(rng: random.Random, _depth: int) -> str:
    return (
        f"<select{_gen_attrs(rng)}><option>{rng.choice(_ELEM_POOL)}</option></select>"
    )


#: 块级发生器权重表
_BLOCK_GENS: list[tuple[float, Callable[[random.Random, int], str]]] = [
    (0.34, _bk_textblock),
    (0.12, _bk_container),
    (0.08, _bk_list),
    (0.06, _bk_table),
    (0.06, _bk_figure),
    (0.05, _bk_details),
    (0.05, _bk_nav),
    (0.05, _bk_pre),
    (0.04, _bk_dl),
    (0.04, _bk_blockquote),
    (0.04, lambda rng, _d: rng.choice(["<hr/>", "<br/>", ""])),
    (0.04, lambda rng, _d: rng.choice(_PROSE_POOL)),  # body 直挂裸文本
    (0.03, _bk_select),
]


def _gen_block(rng: random.Random, depth: int = 0) -> str:
    """块级片段。"""
    if depth <= _MAX_NEST:
        cut = rng.random()
        acc = 0.0
        for weight, gen in _BLOCK_GENS:
            acc += weight
            if cut < acc:
                return gen(rng, depth)
    return _bk_textblock(rng, depth)


#: 随机书各可选面的概率（PLR2004：阈值一律提名）
_P_NO_MIMETYPE = 0.08
_P_DOCTYPE = 0.06
_P_HREF_FRAG = 0.1
_P_SPINE_SKIP = 0.12
_P_NCX = 0.4


def _gen_book(rng: random.Random) -> dict[str, bytes]:
    """随机合法书的成员表。

    仍避开：``../`` manifest href（member 簿记须按 normpath 落位）、
    fragment 文档（``_xhtml`` 恒包裹）、doctype 内部子集（模板只出裸
    ``<!DOCTYPE html>``）、PUA 哨兵形字面（会改变 sent/source 文本关系
    破子串 oracle）、超深嵌套——均由文件尾定向回归钉覆盖。
    """
    opf_dir = rng.choice(["OEBPS", "OEBPS", "OPS", ""])
    opf_path = f"{opf_dir}/content.opf" if opf_dir else "content.opf"
    n_docs = rng.randint(1, 3)
    doc_names = [f"ch{i + 1}.xhtml" for i in range(n_docs)]
    members: dict[str, bytes] = {
        "mimetype": _MIMETYPE,
        _CONTAINER_PATH: _container_xml(opf_path).encode(),
    }
    if rng.random() < _P_NO_MIMETYPE:
        del members["mimetype"]  # 缺 mimetype 的畸形但可翻输入
    decl = rng.choice(["utf-8", "utf-8", "utf-8", "iso", "bom", None])
    head_extra = rng.choice(
        ["", '<meta charset="utf-8"/>', '<link rel="stylesheet" href="s.css"/>']
    )
    doctype = rng.random() < _P_DOCTYPE
    items: list[tuple[str, str, str]] = []
    spine: list[str] = []
    for i, name in enumerate(doc_names):
        body = "".join(_gen_block(rng) for _ in range(rng.randint(1, 7)))
        path = f"{opf_dir}/{name}" if opf_dir else name
        members[path] = _xhtml(body, decl=decl, head=head_extra, doctype=doctype)
        href = f"{name}#frag" if rng.random() < _P_HREF_FRAG else name
        items.append(
            (
                f"c{i}",
                href,
                rng.choice(
                    ["application/xhtml+xml", "application/xhtml+xml", "text/html"]
                ),
            )
        )
        if rng.random() > _P_SPINE_SKIP:  # 偶尔 manifest-only 文档（尾追枚举面）
            spine.append(f"c{i}")
    if not spine:
        spine.append("c0")
    if rng.random() < _P_NCX:
        labels = [rng.choice(_NCX_LABEL_POOL) for _ in range(rng.randint(1, 3))]
        items.append(("ncx", "toc.ncx", "application/x-dtbncx+xml"))
        ncx_name = f"{opf_dir}/toc.ncx" if opf_dir else "toc.ncx"
        members[ncx_name] = _ncx(labels).encode()
    items.append(("img", "i.png", "image/png"))
    members[f"{opf_dir}/i.png" if opf_dir else "i.png"] = b"\x89PNG\r\n\x1a\n"
    items.append(("css", "s.css", "text/css"))
    members[f"{opf_dir}/s.css" if opf_dir else "s.css"] = b"p{margin:0}"
    members[opf_path] = _opf(items, spine).encode()
    return members


# ------------------------------------------------------------------ 校验器/译文器


def _strict_xml(raw: bytes) -> None:
    """严格 XML 解析——不 recover。"""
    etree.fromstring(raw)


def _decl_is_utf8(raw: bytes) -> bool:
    """文档开头 ``<?xml`` decl 存在时其 encoding 值必须是 utf-8（或缺省）。"""
    head = raw.lstrip(b"\xef\xbb\xbf \t\r\n")[:200]
    if not head.startswith(b"<?xml"):
        return True
    decl_end = head.find(b"?>")
    m = _DECL_ENCODING_RE.search(head[: decl_end if decl_end > 0 else None])
    return m is None or m.group(1).lower() == b"utf-8"


def _localname(el: _Element) -> str:
    return el.tag.rsplit("}", 1)[-1] if isinstance(el.tag, str) else ""


_TEXT_SKIP_TAGS = frozenset(
    [
        "script",
        "style",
        "head",
        "title",
        "template",
        "svg",
        "math",
        "rt",
        "rp",
        "rtc",
        "sup",
        "code",
    ]
)
_NOTE_EPUB_TYPES = frozenset(
    ["footnote", "endnote", "rearnote", "note", "footnotes", "endnotes"]
)
_NOTE_ROLES = frozenset(["doc-footnote", "doc-endnote"])
_CSS_DISPLAY_RE = re.compile(r"(?:^|;)\s*display\s*:\s*([\w-]+)", re.IGNORECASE)


def _text_node_skipped(node: NavigableString) -> bool:
    """``epub._ancestor_skip_reason`` 的 oracle 复算——被跳过的文本不进 unit
    text，但会出现在 ``get_text`` 里造成粘连间隙，子串校验必须同样剔除。"""
    hidden = note = False
    for anc in node.parents:
        name = getattr(anc, "name", None)
        if name is None:
            break
        if name in _TEXT_SKIP_TAGS:
            return True
        etype = anc.get("epub:type") or ""
        if "pagebreak" in etype or "page-list" in etype:
            return True
        role = anc.get("role") or ""
        if role == "doc-pagebreak":
            return True
        if any(t in _NOTE_EPUB_TYPES for t in etype.split()) or (role in _NOTE_ROLES):
            note = True
        if not hidden:
            if anc.has_attr("hidden"):
                hidden = True
            else:
                m = _CSS_DISPLAY_RE.search(anc.get("style") or "")
                if m and m.group(1).lower() == "none":
                    hidden = True
    return hidden and not note


def _doc_visible_text(soup: BeautifulSoup) -> str:
    """iter_units 视角的可见文本：只收 ``type is NavigableString``（epub.py:522
    的 exact-type 判定排掉 Comment/PI/CData/Doctype），再剔除被跳子树。"""
    parts = [
        str(n)
        for n in soup.find_all(string=True)
        if type(n) is NavigableString and not _text_node_skipped(n)
    ]
    return normalize_text("".join(parts))


_LONG_ZH_LEN = 300

#: ``_ChaosTranslator`` 译文变体表：``(content, wire tokens) -> reply``。
#: 各变体刻意覆盖 markup/控制字符/echo/丢 token/重复/臆造/空译等对抗面。
_CHAOS_VARIANTS: list[Callable[[str, list[str]], str]] = [
    lambda _c, t: f"这是译文 {' '.join(t)}".strip(),
    lambda _c, t: f"译<>&\"']]>{' '.join(t)}",
    lambda c, _t: c,  # echo → unchanged
    lambda _c, t: f"译\x00\x0b\x1f文{' '.join(t)}",
    lambda _c, t: "长" * _LONG_ZH_LEN + " ".join(t),
    lambda _c, t: " ".join(t),  # 纯 token / 空
    lambda _c, _t: "",  # 空译文（回归钉面——结构断言仍须过）
    lambda _c, t: f"译文 {' '.join(t[1:])}".strip(),  # 丢 token → 阶梯走穿
    lambda _c, t: f"译{' '.join(t)} {t[0] if t else ''}".strip(),  # 重复 token
    lambda _c, t: f"نص {' '.join(t)} 🎉行",
    lambda _c, t: f"译 [[FAKE_7]] {' '.join(t)}".strip(),  # 臆造 token → 走穿
    lambda c, _t: f"{c} 译",  # 近 echo——literals 忠实回显
]


def _protocol_reply(
    user: str,
    response_format: dict[str, str] | None,
    *,
    slot_fn: Callable[[object], str],
    line_fn: Callable[[re.Match[str]], str],
    fallback_fn: Callable[[str], str],
) -> str:
    """MockTranslator 协议骨架——json_object slots / ``[n]`` 编号行 / 裸文回退。

    ``_ChaosTranslator``/``_EmptyTranslator`` 共用；``line_fn`` 独立参化——
    ``_Empty`` 的纯序号桩 ``[n]``（无空格无体）是被钉的 wire 形，不得并成
    统一 ``[n] body`` 前缀。
    """
    if response_format is not None and response_format.get("type") == "json_object":
        try:
            slots = json.loads(user).get("slots") or {}
        except json.JSONDecodeError:
            return "{}"
        return json.dumps({k: slot_fn(v) for k, v in slots.items()}, ensure_ascii=False)
    lines = user.split("\n")
    if lines and all(_NUM_LINE_RE.match(ln) for ln in lines if ln.strip()):
        return "\n".join(
            line_fn(m) if (m := _NUM_LINE_RE.match(ln)) else ln for ln in lines
        )
    return fallback_fn(user)


class _ChaosTranslator:
    """确定性对抗译文器：变体 = md5(content) % N——同输入恒同输出。

    wire 内 token（签发 ``[[X_n]]``/转义 ``[[SL_RAW]]``/字面 ``[[SL]]``）
    逐字抄回即 validator 天然成立；丢弃/重复/臆造变体刻意引爆阶梯走到
    skipped/fault 计数。
    """

    def __init__(self) -> None:
        self.calls: list[str] = []

    def _zh(self, content: str) -> str:
        h = int.from_bytes(
            hashlib.md5(content.encode("utf-8", "surrogatepass")).digest()[:8]  # noqa: S324 -- 确定性散列非密码学
        )
        tokens = ANY_PH_RX.findall(content)
        return _CHAOS_VARIANTS[h % len(_CHAOS_VARIANTS)](content, tokens)

    async def translate(
        self,
        *,
        system: str,  # noqa: ARG002 -- Translator 协议面
        user: str,
        temperature: float,  # noqa: ARG002 -- 同上
        max_tokens: int,  # noqa: ARG002 -- 同上
        response_format: dict[str, str] | None = None,
    ) -> str:
        """MockTranslator 同协议：批行回显编号 / JSON slots 同构。"""
        self.calls.append(user)
        return _protocol_reply(
            user,
            response_format,
            slot_fn=lambda v: self._zh(str(v)),
            line_fn=lambda m: f"{m.group(1)} {self._zh(m.group(2))}",
            fallback_fn=self._zh,
        )


class _EmptyTranslator:
    """每槽位恒空译文——钉"空 zh 仍插仍计 translated"缺陷用。"""

    def __init__(self) -> None:
        self.calls: list[str] = []

    async def translate(
        self,
        *,
        system: str,  # noqa: ARG002
        user: str,
        temperature: float,  # noqa: ARG002
        max_tokens: int,  # noqa: ARG002
        response_format: dict[str, str] | None = None,
    ) -> str:
        self.calls.append(user)
        return _protocol_reply(
            user,
            response_format,
            slot_fn=lambda _v: "",
            line_fn=lambda m: m.group(1),
            fallback_fn=lambda _u: "",
        )


class _AuthFailTranslator:
    """恒抛 401——三连即触 ``AuthTrippedError``，驱动半成品落盘路。"""

    def __init__(self) -> None:
        self.calls: list[str] = []

    async def translate(
        self,
        *,
        system: str,  # noqa: ARG002
        user: str,
        temperature: float,  # noqa: ARG002
        max_tokens: int,  # noqa: ARG002
        response_format: dict[str, str] | None = None,  # noqa: ARG002
    ) -> str:
        self.calls.append(user)
        msg = "401 bad key"
        raise AuthError(msg, status=401)


def _check_epub_out(  # noqa: C901, PLR0912, PLR0915 -- 出包不变量逐条断言，分支即清单
    src: Path,
    dst: Path,
    rep: ExportReport,
    target_lang: str,
    src_members: dict[str, bytes],
) -> None:
    """EPUB 出包全套不变量（见模块 docstring L3 清单）。"""
    assert rep.format == "epub"
    assert sniff_format(dst) == "epub"  # 产出必须可被自家嗅探识别
    with zipfile.ZipFile(dst) as z:
        assert z.testzip() is None
        infos = z.infolist()
        assert infos[0].filename == "mimetype"
        assert infos[0].compress_type == zipfile.ZIP_STORED
        assert z.read("mimetype") == _MIMETYPE
        for info in infos[1:]:
            assert info.compress_type == zipfile.ZIP_DEFLATED
        out_members = {i.filename: z.read(i.filename) for i in infos}
    assert set(out_members) == set(src_members) | {"mimetype"}

    book = load_epub(src)
    soups = _book_soups(book)
    units = list(iter_units(book, soups))
    assert rep.units == len(units)
    assert rep.translated + rep.unchanged + rep.skipped + rep.fault == rep.units
    assert rep.documents == len(book.doc_paths)
    assert all(isinstance(w, str) and w for w in rep.warnings)
    assert not dst.with_name(dst.name + ".state").exists()  # 成功即清理

    for path in book.doc_paths:
        _strict_xml(out_members[path])
        assert _decl_is_utf8(out_members[path]), path
    _strict_xml(out_members[book.opf_path])

    opf_root = etree.fromstring(out_members[book.opf_path])
    langs = [el.text for el in opf_root.iter() if el.tag == _DC_LANG_TAG]
    expected_lang = safe_language(target_lang) or "en"
    assert langs == [expected_lang]

    ncx_applied = 0
    if book.ncx_path and book.ncx_path in out_members:
        out_ncx = etree.fromstring(out_members[book.ncx_path])
        in_ncx = etree.fromstring(src_members[book.ncx_path])
        out_texts = [el for el in out_ncx.iter() if _localname(el) == "text"]
        in_texts = [el for el in in_ncx.iter() if _localname(el) == "text"]
        assert len(out_texts) == len(in_texts)
        for out_el, in_el in zip(out_texts, in_texts, strict=True):
            orig = normalize_text("".join(in_el.itertext()))
            if not orig or is_special_text(orig) or is_apparatus_text(orig):
                assert out_el.text == in_el.text  # 未枚举的原样保留
                continue
            new = out_el.text or ""
            if new == orig:
                continue  # echo/占位短路未改写
            assert new.startswith(orig + " / "), new[:80]
            ncx_applied += 1

    zh_total = 0
    for path in book.doc_paths:
        out_soup = BeautifulSoup(out_members[path], "html.parser")
        zh_nodes = out_soup.select(".texlate-zh")
        zh_total += len(zh_nodes)
        doc_units = [u for u in units if u.soup is soups[path]]
        literals: set[str] = set()
        for u in doc_units:
            literals |= set(find_markers(u.text)) - set(u.markers)
        zh_text = "".join(n.get_text() for n in zh_nodes)
        for tok in find_markers(zh_text):
            assert tok in literals, (
                f"译文节点残留非字面 token {tok}（签发未还原/臆造未剥）"
            )
        doc_text = _doc_visible_text(out_soup)
        for u in doc_units:
            for kind, seg in split_on_markers(u.text, find_markers(u.text)):
                if kind != "text":
                    continue
                seg_n = normalize_text(seg)
                if seg_n:
                    assert seg_n in doc_text, u.job_id  # 源文未被吃掉
    assert zh_total + ncx_applied == rep.translated


# ------------------------------------------------------------------ L1 纯函数


#: reconcile 回复部件组成概率阈值
_CUT_REPLY_ISSUED = 0.35
_CUT_REPLY_LIT = 0.55
_CUT_REPLY_FAKE = 0.7


def test_fuzz_reconcile_markers() -> None:  # noqa: C901 -- 对抗回复生成即分支表
    """调和协议：签发恰一次、字面保留、臆造剥净、一致回复逐字节原样。"""
    rng = fuzz_rng(20260917)
    for _ in range(_ITERS_PURE):
        issued = [
            marker_token(rng.choice(["IMG", "CODE", "A"]), i + 1)
            for i in range(rng.randint(0, 4))
        ]
        lit_pool = [f"[[LIT_{i}]]" for i in range(rng.randint(0, 3))]
        sent = " ".join(
            rng.choice(["word", "文字", *issued, *lit_pool])
            for _ in range(rng.randint(1, 10))
        )
        sent_literal = set(find_markers(sent)) - set(issued)  # 源文自有字面
        parts: list[str] = []
        for _ in range(rng.randint(0, 8)):
            r = rng.random()
            if r < _CUT_REPLY_ISSUED and issued:
                parts.append(rng.choice(issued) * rng.choice([1, 1, 2]))  # 偶发重复
            elif r < _CUT_REPLY_LIT and lit_pool:
                parts.append(rng.choice(lit_pool))
            elif r < _CUT_REPLY_FAKE:
                parts.append(f"[[FAKE_{rng.randrange(99)}]]")
            else:
                parts.append(rng.choice(["译文", "tail", "x"]))
        reply = " ".join(parts)
        out = reconcile_markers(sent, reply, issued=issued)
        for tok in issued:
            assert out.count(tok) == 1  # 签发 token 恰一次（重复剥 + 丢补尾）
        for tok in sent_literal:
            assert out.count(tok) == reply.count(tok)  # 字面逐处保留
        for m in MARKER_RE.finditer(reply):
            tok = m.group(0)
            if tok not in issued and tok not in sent_literal:
                assert tok not in out  # 臆造剥净
        reply_tokens = MARKER_RE.findall(reply)
        if (
            not set(issued) - set(reply_tokens)
            and all(reply.count(t) == 1 for t in issued)
            and all(t in issued or t in sent_literal for t in reply_tokens)
        ):
            assert out == reply  # 已一致回复逐字节原样


def test_fuzz_split_on_markers_roundtrip() -> None:
    """``split_on_markers`` 恒还原输入；marker 段值 ∈ tokens。"""
    rng = fuzz_rng(20260918)
    for _ in range(_ITERS_PURE):
        tokens = [f"[[T_{i}]]" for i in range(rng.randint(0, 5))]
        text = "".join(
            rng.choice(["w", "中", *tokens, "[[X_9]]", " ", "[["])
            for _ in range(rng.randint(0, 12))
        )
        pieces = split_on_markers(text, tokens)
        assert "".join(v for _k, v in pieces) == text
        for kind, value in pieces:
            if kind == "marker":
                assert value in tokens


def test_fuzz_normalize_sanitize() -> None:
    """归一化/净化幂等；输出无 XML 非法字符、零宽字符、连续/首尾空白。"""
    rng = fuzz_rng(20260919)
    pool = "abc中 \t\n\x00\x0b\x1f\u00ad\u200b\ufeff[]_XY9🎉<>"
    for _ in range(_ITERS_PURE):
        s = "".join(rng.choice(pool) for _ in range(rng.randint(0, 60)))
        once = normalize_text(s)
        assert normalize_text(once) == once
        assert not _XML_ILLEGAL_CHARS.search(once)
        assert not any(c in once for c in _INVISIBLE)
        assert "  " not in once
        assert once == once.strip()
        clean = sanitize_xml_text(s)
        assert sanitize_xml_text(clean) == clean
        assert not _XML_ILLEGAL_CHARS.search(clean)


def test_fuzz_safe_language() -> None:
    """任意输入不炸；非 None 输出恒 ``[\\w-]+`` fullmatch。"""
    rng = fuzz_rng(20260920)
    pool = "abzhCN-_ <>\"'&\n\x00中🎉/\\;"
    assert safe_language(None) is None
    assert safe_language("") is None
    for _ in range(_ITERS_PURE):
        s = "".join(rng.choice(pool) for _ in range(rng.randint(0, 12)))
        out = safe_language(s)
        assert out is None or re.fullmatch(r"[\w-]+", out)


def test_fuzz_ordinals_and_marker_names() -> None:
    """``Ordinals.allocate`` 永不发占用 token；``marker_name`` 恒 ``[A-Z_]+``。"""
    rng = fuzz_rng(20260921)
    # find_markers 首现序去重（模块 docstring 声称的不变量——确定性钉在 fuzz 环外）
    assert find_markers("[[ZZ_9]] [[ZZ_9]]") == ["[[ZZ_9]]"]
    assert find_markers("a [[ZZ_9]] b [[X_1]] c [[ZZ_9]]") == [
        "[[ZZ_9]]",
        "[[X_1]]",
    ]
    for _ in range(_ITERS_PURE):
        ords = Ordinals(start=rng.randint(0, 3))
        occupied = (
            " ".join(f"[[IMG_{i}]]" for i in range(rng.randint(0, 4)))
            + " prose [[A_1]]"
        )
        taken: list[str] = []
        for _k in range(rng.randint(1, 6)):
            tok = ords.allocate(
                rng.choice(["img", "code", "x-y", None]),
                occupied=occupied,
                taken=taken,
            )
            assert tok not in occupied
            assert tok not in taken
            assert MARKER_RE.fullmatch(tok)
            taken.append(tok)
        tag = "".join(rng.choice("aZ09_-x ") for _ in range(rng.randint(0, 8)))
        assert re.fullmatch(r"[A-Z_]+", marker_name(tag))
        for m in find_markers(occupied + " [[ZZ_9]] [[ZZ_9]]"):
            assert MARKER_RE.fullmatch(m)


def test_fuzz_text_predicates_never_raise() -> None:
    """过滤/判定谓词任意对抗串不炸、返回 bool。"""
    rng = fuzz_rng(20260922)
    for _ in range(_ITERS_PURE):
        s = "".join(
            rng.choice("a中1 .,!?:;https://x [[T_1]]\x00­")
            for _ in range(rng.randint(0, 50))
        )
        for pred in (
            is_special_text,
            is_apparatus_text,
            is_wordless,
            is_placeholder_only,
        ):
            assert isinstance(pred(s), bool)


def test_coerce_glossary_edges(tmp_path: Path) -> None:
    """``coerce_glossary``：mapping → ``Glossary``；缺席路径 → ``ExportError``。"""
    g = coerce_glossary({"alpha": "阿尔法", "  ": "x", "beta": ""})
    assert g is not None
    assert coerce_glossary(None) is None
    with pytest.raises(ExportError):
        coerce_glossary(tmp_path / "missing.yaml")


# ------------------------------------------------------------------ L2 单元级

_HOSTILE_BODIES = [
    "<p>Plain para <code>c0de</code> tail.</p>",
    '<div>pre text <img src="i.png"/> mid <p>inner block</p> post tail</div>',
    "<figure><img src='i.png'/><figcaption>cap text here</figcaption></figure>",
    (
        '<nav epub:type="doc-toc"><ol><li><a href="c.xhtml">Ch One</a></li>'
        "<li><span>Ch Two</span></li></ol></nav>"
    ),
    "<ul><li>item one <sup>2</sup> rest</li><li>item two</li></ul>",
    "<table><tr><td>cell a <em>em</em></td><td>cell b</td></tr></table>",
    "<p>ruby <ruby>漢<rt>kan</rt></ruby> tail here.</p>",
    "<p><a href='x.xhtml'>whole link run covering all</a></p>",
    "<p><b><i><u>deep inline cover</u></i></b> more text</p>",
    "<p><span>short</span></p>",
    "<p>ends with marker <code>only code</code></p>",
    "<p>marker first <img src='i.png'/> then prose words follow.</p>",
    "<h2>Title <sup>1</sup> heading</h2><p>body words after heading.</p>",
    "<blockquote><p>quoted prose with <em>emphasis</em> tail.</p></blockquote>",
    "<p>literal [[IMG_1]] plus <img src='i.png'/> real marker tail.</p>",
    "<p>bare [[SL]] and [[PL]] literals sit among prose words here.</p>",
    "<details><summary>sum text</summary><p>det body</p></details>",
    "<dl><dt>term def</dt><dd>defn body text</dd></dl>",
    "<p><svg><text>SVGTXT</text></svg> tail words after svg.</p>",
    "<p><iframe>iframe fallback text</iframe> after words.</p>",
    "<p hidden>hidden para</p><p>visible para words here.</p>",
    '<p><span epub:type="pagebreak">5</span>page text continues.</p>',
    "<p>alpha<br/>beta split lines.</p>",
    "<p>Text with \x00control\x0b chars embedded.</p>",
    "<p>   </p><p>second para real words.</p>",
    "<p>Figure 3</p><p>Real prose paragraph words.</p>",
    "<p>https://example.com/only</p><p>More real prose words.</p>",
    "<select><option>opt one</option><option>opt two</option></select>",
    "<p><math><mtext>math text</mtext></math> tail words.</p>",
    "<p>Comment tail <!-- inline comment --> after.</p>",
    "<p><del>deleted words</del><ins>inserted words here</ins> tail.</p>",
    "<p dir='rtl'>نص عربي كامل هنا للاختبار</p>",
    "<div><p>nested one</p><div><p>nested two deep</p></div></div>",
]

_ZH_PARTS = [
    "译文文本",
    "",
    "   ",
    "带<标记>&\"'字符",
    "tail]]>cdata",
    "ctl\x00\x0b\x01char",
    "换行\n两行",
    "🎉🚀 emoji",
    "长" * 200,
    "5 pua",
    "[[FAKE_7]]",
    "[[IMG_99]]",
    "[[INVENT_2]]",
    "orig [[SL]]",
    "普通 zh 句。",
]


#: ``_chaos_zh`` 部件组成概率阈值
_CUT_ZH_ISSUED = 0.3
_CUT_ZH_PART = 0.55
_CUT_ZH_ECHO = 0.65


def _chaos_zh(rng: random.Random, issued: list[str], text: str) -> str:
    parts: list[str] = []
    for _ in range(rng.randint(1, 4)):
        r = rng.random()
        if r < _CUT_ZH_ISSUED and issued:
            parts.append(rng.choice(issued) * rng.choice([1, 1, 2]))  # 偶发重复
        elif r < _CUT_ZH_PART:
            parts.append(rng.choice(_ZH_PARTS))
        elif r < _CUT_ZH_ECHO:
            parts.append(text)  # echo
        else:
            parts.append("译" + str(rng.randint(0, 999)))
    return " ".join(parts)


def _assert_unit_shape(units: list[Unit]) -> None:
    """``iter_units`` 产出形态断言集（hostile/insert 两路共用）。"""
    ids: set[str] = set()
    for u in units:
        assert u.job_id not in ids
        ids.add(u.job_id)
        assert u.text
        assert u.text == normalize_text(u.text)
        assert not is_special_text(u.text)
        assert not is_apparatus_text(u.text)
        if u.ncx_text is None:
            # NCX 枚举缺 is_placeholder_only 过滤（已观察，未钉）
            assert not is_placeholder_only(u.text)
            for tok in u.markers:
                assert MARKER_RE.fullmatch(tok)
                assert u.text.count(tok) == 1  # 签发 token 在 sent 恰一次


def test_fuzz_iter_units_hostile(tmp_path: Path) -> None:
    """``iter_units`` 对 hostile body 不抛；unit 字段自洽。"""
    rng = fuzz_rng(20260923)
    for i in range(_ITERS_UNITS):
        body = (
            _HOSTILE_BODIES[i]
            if i < len(_HOSTILE_BODIES)
            else "".join(_gen_block(rng) for _ in range(rng.randint(1, 6)))
        )
        ncx = (
            [rng.choice(_NCX_LABEL_POOL) for _ in range(rng.randint(1, 3))]
            if rng.random() < _P_NCX
            else None
        )
        src = _write_epub(tmp_path, f"u{i}", _mini_book(body, ncx_labels=ncx))
        book = load_epub(src)
        soups = _book_soups(book)
        units = list(iter_units(book, soups))  # 不得抛
        _assert_unit_shape(units)


def test_fuzz_insert_translation_chaos(tmp_path: Path) -> None:
    """``insert_translation`` 对抗译文不抛；插后 well-formed；token 协议成立。"""
    rng = fuzz_rng(20260924)
    for i in range(_ITERS_INSERT):
        body = (
            _HOSTILE_BODIES[i]
            if i < len(_HOSTILE_BODIES)
            else "".join(_gen_block(rng) for _ in range(rng.randint(1, 6)))
        )
        ncx = (
            [rng.choice(_NCX_LABEL_POOL) for _ in range(rng.randint(1, 3))]
            if rng.random() < _P_NCX
            else None
        )
        src = _write_epub(tmp_path, f"i{i}", _mini_book(body, ncx_labels=ncx))
        book = load_epub(src)
        soups = _book_soups(book)
        units = list(iter_units(book, soups))
        for u in units:
            zh = _chaos_zh(rng, list(u.markers), u.text)
            warn = insert_translation(u, zh, _ZH_LANG)  # 不得抛
            assert warn is None or isinstance(warn, str)
            if u.ncx_text is not None:
                expected_zh = sanitize_xml_text(
                    reconcile_markers(u.text, zh, issued=u.markers)
                )
                assert u.ncx_text.text == f"{u.text} / {expected_zh}"
        for soup in soups.values():
            raw = sanitize_xml_text(soup.encode("utf-8").decode("utf-8")).encode()
            _strict_xml(raw)
            doc_units = [u for u in units if u.soup is soup]
            literals: set[str] = set()
            for u in doc_units:
                literals |= set(find_markers(u.text)) - set(u.markers)
            zh_text = "".join(n.get_text() for n in soup.select(".texlate-zh"))
            for tok in find_markers(zh_text):
                assert tok in literals, f"签发/臆造 token 残留: {tok}"


def test_fuzz_docx_insert_after() -> None:
    """``insert_after`` 对抗译文不抛；克隆段紧邻 + 色戳/语言章/剥 paraId。"""
    rng = fuzz_rng(20260925)
    doc = Document()
    for i in range(30):
        doc.add_paragraph(f"Source para {i} {rng.randrange(10**6)} text.")
    paras = list(doc.element.body.iter(qn("w:p")))  # 物化——插译会改树
    for i, p_el in enumerate(paras):
        if i % _PARAID_EVERY == 0:
            p_el.set(qn("w14:paraId"), f"{i:08X}")  # 克隆必须剥的唯一锚
        zh = _chaos_zh(rng, [], Paragraph(p_el, None).text)
        insert_after(p_el, zh, _ZH_LANG)
        clone = p_el.getnext()
        assert clone is not None
        assert clone.tag == qn("w:p")
        assert clone.get(qn("w14:paraId")) is None
        assert clone.get(qn("w14:textId")) is None
        assert {c.tag for c in clone} <= {qn("w:pPr"), qn("w:r")}
        assert clone.find(f".//{qn('w:sectPr')}") is None
        colors = [c.get(qn("w:val")) for c in clone.iter(qn("w:color"))]
        assert "555555" in colors
        langs = [el.get(qn("w:val")) for el in clone.iter(qn("w:lang"))]
        assert _ZH_LANG in langs
        assert Paragraph(clone, None).text == sanitize_xml_text(zh)
    _strict_xml(etree.tostring(doc.element))


# ------------------------------------------------------------------ L3 端到端


def test_fuzz_epub_e2e(tmp_path: Path) -> None:
    """随机合法书 × 对抗译文器：出包全套不变量。"""
    rng = fuzz_rng(20260926)
    for i in range(_ITERS_EPUB):
        members = _gen_book(rng)
        src = _write_epub(tmp_path, f"e{i}", members)
        dst = tmp_path / f"e{i}.out.epub"
        lang = rng.choice([_ZH_LANG, _ZH_LANG, _ZH_LANG, "zh<x>", "en"])
        seen: list[ChunkResult] = []
        rep = translate_epub(
            src,
            dst,
            _ChaosTranslator(),
            target_lang=lang,
            on_result=seen.append,
        )
        assert len(seen) == rep.units  # 每单元恰好回调一次
        assert len({r.chunk_id for r in seen}) == rep.units
        _check_epub_out(src, dst, rep, lang, members)


def test_fuzz_docx_e2e(tmp_path: Path) -> None:
    """随机 docx × 对抗译文器：源段子序列 + 紧邻 + 色戳 + 全 part XML 合法。"""
    rng = fuzz_rng(20260927)
    for i in range(_ITERS_DOCX):
        src = tmp_path / f"d{i}.docx"
        src.write_bytes(_gen_docx(rng))
        dst = tmp_path / f"d{i}.out.docx"
        lang = rng.choice([_ZH_LANG, _ZH_LANG, "zh<x>", "en"])
        seen: list[ChunkResult] = []
        rep = translate_docx(
            src,
            dst,
            _ChaosTranslator(),
            target_lang=lang,
            on_result=seen.append,
        )
        assert len(seen) == rep.units
        _check_docx_out(src, dst, rep, lang)


# ------------------------------------------------------------------ DOCX 构造/校验

_DOCX_EXTRA_POOL = [
    "docx [[IMG_1]] literal",
    "tab\there\nnewline",
    "w:noBreakHyphen neighbor",
]


def _docx_text_pool(rng: random.Random) -> str:
    """DOCX 文本池——剥 XML 非法字符（lxml 写入侧本身就拒）。"""
    return sanitize_xml_text(rng.choice(_PROSE_POOL + _DOCX_EXTRA_POOL))


def _inject_raw(p_el: _Element, xml: str) -> None:
    """``w:p`` 尾部追加原始 OOXML 片段（fldSimple/del/instrText/oMath/sdt）。"""
    p_el.append(parse_xml(xml))


#: ``_gen_docx`` 各对抗面概率（PLR2004：阈值一律提名）
_P_DOCX_HEADER = 0.35
_P_DOCX_TABLE = 0.12
_P_DOCX_CELL_EXTRA = 0.4
_P_DOCX_STYLE = 0.2
_P_DOCX_TOC = 0.12
_P_DOCX_BOLD = 0.15
_P_DOCX_HIDDEN = 0.1
_P_DOCX_INSTR = 0.1
_P_DOCX_FLD = 0.08
_P_DOCX_DEL = 0.08
_P_DOCX_OMATH = 0.06
_P_DOCX_NBH = 0.08
_P_DOCX_SDT = 0.06
_P_DOCX_FOOTNOTES = 0.35


def _gen_docx(rng: random.Random) -> bytes:  # noqa: C901, PLR0912 -- 注入面枚举即分支表
    """python-docx 随机书 + header part + zip 手术注入 footnotes。"""
    doc = Document()
    if rng.random() < _P_DOCX_HEADER:
        hdr = doc.sections[0].header
        hdr.is_linked_to_previous = False  # 不先解链写不进 header part
        hdr.paragraphs[0].text = _docx_text_pool(rng)
    for _ in range(rng.randint(2, 14)):
        r = rng.random()
        if r < _P_DOCX_TABLE:
            tbl = doc.add_table(rows=rng.randint(1, 2), cols=rng.randint(1, 2))
            for row in tbl.rows:
                for cell in row.cells:
                    cell.paragraphs[0].add_run(_docx_text_pool(rng))
                    if rng.random() < _P_DOCX_CELL_EXTRA:
                        cell.add_paragraph(_docx_text_pool(rng))
            continue
        p = doc.add_paragraph()
        if rng.random() < _P_DOCX_STYLE:
            with suppress(KeyError):
                p.style = doc.styles[
                    rng.choice(["Heading 1", "Heading 2", "Title", "Quote"])
                ]
        if rng.random() < _P_DOCX_TOC:  # TOC 样式 → 跳过枚举
            ppr = p._p.get_or_add_pPr()  # noqa: SLF001 -- oxml 面
            st = ppr.get_or_add_pStyle()
            st.set(qn("w:val"), rng.choice(["TOC1", "toc 2", "Contents"]))
        for _ in range(rng.randint(1, 4)):
            run = p.add_run(_docx_text_pool(rng) + " ")
            if rng.random() < _P_DOCX_BOLD:
                run.font.bold = True
            if rng.random() < _P_DOCX_HIDDEN:
                run.font.hidden = True  # w:vanish → 隐藏 run 不送模型
        if rng.random() < _P_DOCX_INSTR:
            _inject_raw(
                p._p,  # noqa: SLF001
                f"<w:r {nsdecls('w')}><w:instrText> TOC \\o </w:instrText></w:r>",
            )
        if rng.random() < _P_DOCX_FLD:
            _inject_raw(
                p._p,  # noqa: SLF001
                f'<w:fldSimple {nsdecls("w")} w:instr=" X ">'
                f"<w:r><w:t>fldtext</w:t></w:r></w:fldSimple>",
            )
        if rng.random() < _P_DOCX_DEL:
            _inject_raw(
                p._p,  # noqa: SLF001
                f"<w:del {nsdecls('w')}><w:r><w:delText>gone</w:delText></w:r></w:del>",
            )
        if rng.random() < _P_DOCX_OMATH:
            _inject_raw(
                p._p,  # noqa: SLF001
                f"<m:oMath {nsdecls('w', 'm')}><m:r><m:t>eq</m:t></m:r></m:oMath>",
            )
        if rng.random() < _P_DOCX_NBH:
            _inject_raw(
                p._p,  # noqa: SLF001
                f"<w:r {nsdecls('w')}><w:t>a</w:t><w:noBreakHyphen/><w:t>b</w:t></w:r>",
            )
        if rng.random() < _P_DOCX_SDT:
            _inject_raw(
                p._p,  # noqa: SLF001
                f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
                f"<w:t>nested p text</w:t></w:r></w:p></w:sdtContent></w:sdt>",
            )
    buf = io.BytesIO()
    doc.save(buf)
    blob = buf.getvalue()
    if rng.random() < _P_DOCX_FOOTNOTES:
        blob = _docx_inject_footnotes(blob, rng)
    return blob


def _docx_inject_footnotes(blob: bytes, rng: random.Random) -> bytes:
    """zip 手术：footnotes part + content-type override + document rel。"""
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    paras = "".join(
        f"<w:p><w:r><w:t>{escape(_docx_text_pool(rng))}</w:t></w:r></w:p>"
        for _ in range(rng.randint(1, 3))
    )
    members["word/footnotes.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<w:footnotes {nsdecls('w')}>"
        f'<w:footnote w:id="1">{paras}</w:footnote></w:footnotes>'
    ).encode()
    ct = members["[Content_Types].xml"].decode()
    members["[Content_Types].xml"] = ct.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    ).encode()
    rels = members["word/_rels/document.xml.rels"].decode()
    members["word/_rels/document.xml.rels"] = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
        "</Relationships>",
    ).encode()
    return zip_bytes(members.items())


def _is_zh_para(p_el: _Element) -> bool:
    # 只看直系 w:r——嵌套 w:p（sdt 内）里的 zh 克隆会污染子孙遍历
    return any(
        c.get(qn("w:val")) == "555555"
        for r in p_el.findall(qn("w:r"))
        for c in r.iter(qn("w:color"))
    )


def _para_text_oracle(p_el: _Element) -> str:
    """``docx._para_text`` 的 oracle 复算：w:t + 字面节点 + 保护子树跳过。"""
    skip = {qn("w:instrText"), qn("w:delText")}
    lit = {
        qn("w:tab"): "\t",
        qn("w:br"): "\n",
        qn("w:cr"): "\n",
        qn("w:noBreakHyphen"): "-",
    }
    protected = {
        qn("m:oMath"),
        qn("m:oMathPara"),
        qn("w:fldSimple"),
        qn("w:del"),
        qn("w:moveFrom"),
        qn("w:rt"),
        qn("w:p"),
    }
    hidden = (qn("w:vanish"), qn("w:specVanish"))
    parts: list[str] = []
    for node in p_el.iter():
        anc = node.getparent()
        prot = False
        while anc is not None and anc is not p_el:
            if anc.tag in protected:
                prot = True
                break
            if anc.tag == qn("w:r"):
                rpr = anc.find(qn("w:rPr"))
                # CT_OnOff 镜像 docx._protected：w:val 缺省=开，0/false/off 显式关
                if rpr is not None and any(
                    (prop_el := rpr.find(h)) is not None
                    and (prop_el.get(qn("w:val")) or "true").lower()
                    not in ("0", "false", "off")
                    for h in hidden
                ):
                    prot = True
                    break
            anc = anc.getparent()
        if prot:
            continue
        if node.tag in skip:
            continue
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag in lit:
            parts.append(lit[node.tag])
    return normalize_text("".join(parts))


def _check_docx_out(  # noqa: C901 -- 出包不变量逐条断言，分支即清单
    src: Path, dst: Path, rep: ExportReport, target_lang: str
) -> None:
    """DOCX 出包全套不变量（见模块 docstring L3 清单）。"""
    assert rep.format == "docx"
    assert sniff_format(dst) == "docx"
    with zipfile.ZipFile(dst) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                _strict_xml(z.read(name))
    assert not dst.with_name(dst.name + ".state").exists()

    src_doc = Document(str(src))
    src_pairs = list(iter_docx_units(src_doc))
    assert rep.units == len(src_pairs)
    assert rep.translated + rep.unchanged + rep.skipped + rep.fault == rep.units
    extra_parts = {u.part_name for u, part, _root in src_pairs if part is not None}
    assert rep.documents == 1 + len(extra_parts)

    out_doc = Document(str(dst))
    src_paras = list(src_doc.element.body.iter(qn("w:p")))
    out_paras = list(out_doc.element.body.iter(qn("w:p")))
    non_zh = [p for p in out_paras if not _is_zh_para(p)]
    # 源段零丢失且原序保持（译文段是新增克隆）
    assert len(non_zh) == len(src_paras)
    for got, want in zip(non_zh, src_paras, strict=True):
        assert _para_text_oracle(got) == _para_text_oracle(want)
    # 译文段是源段的紧邻兄弟（addnext 语义；DFS 序里嵌套 p 会插队）
    for p in out_paras:
        if _is_zh_para(p):
            prev = p.getprevious()
            assert prev is not None
            assert prev.tag == qn("w:p")
            assert not _is_zh_para(prev)
    # 全 part 译文段总数 == translated；safe lang 时逐段语言章
    safe = safe_language(target_lang)
    zh_total = 0
    with zipfile.ZipFile(dst) as z:
        for name in z.namelist():
            if not _DOCX_TEXT_PART_RE.fullmatch(name):
                continue
            root = etree.fromstring(z.read(name))
            for p in root.iter(qn("w:p")):
                if not _is_zh_para(p):
                    continue
                zh_total += 1
                if safe is not None:
                    vals = [el.get(qn("w:val")) for el in p.iter(qn("w:lang"))]
                    assert safe in vals, name
    assert zh_total == rep.translated
    if safe is not None:
        assert out_doc.core_properties.language == safe


# ------------------------------------------------------------------ 半成品/分派


def test_partial_save_on_auth_trip(tmp_path: Path) -> None:
    """auth 熔断 → ``AuthTrippedError``；半成品仍是合法可 sniff 的出包。"""
    body = "".join(
        f"<p>Paragraph number {i} with enough prose words.</p>" for i in range(6)
    )
    src = _write_epub(tmp_path, "auth", _mini_book(body))
    dst = tmp_path / "auth.out.epub"
    with pytest.raises(AuthTrippedError):
        translate_epub(src, dst, _AuthFailTranslator())
    assert sniff_format(dst) == "epub"
    with zipfile.ZipFile(dst) as z:
        assert z.testzip() is None
        assert z.infolist()[0].filename == "mimetype"
        assert z.read("mimetype") == _MIMETYPE
        _strict_xml(z.read("OEBPS/ch1.xhtml"))
        _strict_xml(z.read("OEBPS/content.opf"))
    assert dst.with_name(dst.name + ".state").exists()  # 断点现场保留

    doc = Document()
    for i in range(5):
        doc.add_paragraph(f"Docx para {i} prose words here.")
    dsrc = tmp_path / "auth.docx"
    doc.save(str(dsrc))
    ddst = tmp_path / "auth.out.docx"
    with pytest.raises(AuthTrippedError):
        translate_docx(dsrc, ddst, _AuthFailTranslator())
    with zipfile.ZipFile(ddst) as z:
        assert z.testzip() is None
        _strict_xml(z.read("word/document.xml"))
    assert ddst.with_name(ddst.name + ".state").exists()


def test_on_result_exception_swallowed(tmp_path: Path) -> None:
    """``on_result`` 回调抛错被编排吞掉——导出照常完成（pipeline._collect 护栏）。"""

    def boom(_r: ChunkResult) -> None:
        msg = "callback boom"
        raise RuntimeError(msg)

    src = _write_epub(
        tmp_path, "cb", _mini_book("<p>Para words for callback test.</p>")
    )
    dst = tmp_path / "cb.out.epub"
    rep = translate_epub(src, dst, MockTranslator(), on_result=boom)
    assert rep.translated == rep.units == 1
    assert sniff_format(dst) == "epub"


def test_export_document_dispatch(tmp_path: Path) -> None:
    """内容嗅探分派：docx 身 epub 名 → docx 管线；epub 身 docx 名同理。"""
    doc = Document()
    doc.add_paragraph("Dispatch para with prose words.")
    fake_epub = tmp_path / "fake.epub"
    doc.save(str(fake_epub))
    rep = export_document(fake_epub, tmp_path / "o.docx", _ChaosTranslator())
    assert rep.format == "docx"

    real_epub = _write_epub(
        tmp_path,
        "real",
        _mini_book("<p>Dispatch epub prose words here.</p>"),
    )
    fake_docx = tmp_path / "fake.docx"
    fake_docx.write_bytes(real_epub.read_bytes())
    rep2 = export_document(fake_docx, tmp_path / "o2.epub", _ChaosTranslator())
    assert rep2.format == "epub"

    bad = tmp_path / "bad.epub"
    bad.write_bytes(b"not a zip at all")
    with pytest.raises(UnsupportedFormatError):
        export_document(bad, tmp_path / "o3.epub", _ChaosTranslator())


def test_fuzz_sniff_format(tmp_path: Path) -> None:
    """``sniff_format``：任意成员集不炸；epub/docx 判定与嗅探规则互证。"""
    rng = fuzz_rng(20260928)
    pool = [
        "mimetype",
        "word/document.xml",
        "META-INF/container.xml",
        "OEBPS/c.opf",
        "x",
        "word/footnotes.xml",
    ]
    for i in range(_ITERS_SNIFF):
        members = {
            n: (
                b"application/epub+zip"
                if rng.random() < _P_SNIFF_MIME
                else rng.randbytes(30)
            )
            for n in rng.sample(pool, rng.randint(0, len(pool)))
        }
        p = _write_epub(tmp_path, f"s{i}", members)
        out = sniff_format(p)
        valid_mt = members.get("mimetype", b"")[:64].strip() == _MIMETYPE
        if out == "epub":
            assert valid_mt
        elif out == "docx":
            assert "word/document.xml" in members
            assert not valid_mt  # 否则 epub 判定先命中
        else:
            assert out is None
            assert not valid_mt
            assert "word/document.xml" not in members


# ------------------------------------------------------------- 定向缺陷回归


def _pin_book(doc: str | bytes, tmp: Path, name: str) -> Path:
    """单文档回归钉书：str → ``_xhtml`` 包裹；bytes → 全档逐字节。"""
    ch1 = doc if isinstance(doc, bytes) else _xhtml(doc)
    return _write_epub(tmp, name, _single_doc_book(ch1))


@pytest.mark.parametrize(
    "frag",
    [
        "<script>if (a < b && c) { go(); }</script>",
        "<style>a > b && c { x: 1 }</style>",
    ],
    ids=["script_lt", "style_gt"],
)
def test_rawtext_script_style_escaped(frag: str, tmp_path: Path) -> None:
    """raw-text 元素内的 ``<``/``&`` 预转义 → 成员 well-formed。"""
    src = _pin_book(f"<p>Para text words here.</p>{frag}", tmp_path, "r")
    dst = tmp_path / "r.out.epub"
    translate_epub(src, dst, MockTranslator())
    with zipfile.ZipFile(dst) as z:
        _strict_xml(z.read("OEBPS/ch1.xhtml"))


_VERBATIM_DOCTYPE = (
    b'<?xml version="1.0"?>'
    b'<!DOCTYPE html [<!ENTITY x "y">]>'
    b'<html xmlns="http://www.w3.org/1999/xhtml">'
    b"<head><title>t</title></head>"
    b"<body><p>Para words here.</p></body></html>"
)


@pytest.mark.parametrize(
    "doc",
    [
        "<p>Para words here.</p><!-- tricky -- comment -->",
        _VERBATIM_DOCTYPE,
        "<p>Para words here.</p><?proc <inside> stuff?>",
        "<p>Para words here.</p><p<div>weird</p<div>",
    ],
    ids=["comment_dash", "doctype_subset", "pi_lt", "bad_tagname"],
)
def test_verbatim_constructs_sanitized(doc: str | bytes, tmp_path: Path) -> None:
    """宽容解析构造经 ``_sanitize_dom`` 整形 → 输出成员 strict-parse。"""
    src = _pin_book(doc, tmp_path, "v")
    dst = tmp_path / "v.out.epub"
    translate_epub(src, dst, MockTranslator())
    with zipfile.ZipFile(dst) as z:
        _strict_xml(z.read("OEBPS/ch1.xhtml"))


def test_dotdot_href(tmp_path: Path) -> None:
    """manifest ``../`` href：合法 URI 相对引用须 normpath 解析。"""
    members = {
        "mimetype": _MIMETYPE,
        _CONTAINER_PATH: _container_xml("OEBPS/content.opf").encode(),
        "OEBPS/content.opf": _opf(
            [("c0", "../ch1.xhtml", "application/xhtml+xml")], ["c0"]
        ).encode(),
        "ch1.xhtml": _xhtml("<p>Dotdot chapter prose words here.</p>"),
    }
    src = _write_epub(tmp_path, "dd", members)
    book = load_epub(src)
    assert "ch1.xhtml" in book.doc_paths


def test_pua_sentinel_collision(tmp_path: Path) -> None:
    """源文逐字印着 ``0`` → 不得抢占哨兵位/漏进 sent 文本。"""
    body = '<p>Before 0 literal then <img src="i.png"/> after.</p>'
    src = _pin_book(body, tmp_path, "pua")
    book = load_epub(src)
    units = list(iter_units(book, _book_soups(book)))
    assert units
    for u in units:
        assert "" not in u.text, u.text
        assert "" not in u.text, u.text


def test_fragment_single_root(tmp_path: Path) -> None:
    """单根 fragment 文档 ``<p>x</p>``：插译锚定不造双根。"""
    src = _pin_book(b"<p>Fragment para prose words.</p>", tmp_path, "frag")
    dst = tmp_path / "frag.out.epub"
    translate_epub(src, dst, MockTranslator())
    with zipfile.ZipFile(dst) as z:
        _strict_xml(z.read("OEBPS/ch1.xhtml"))


def test_empty_translation_skipped(tmp_path: Path) -> None:
    """空/纯序号桩译文 → 不插 ``texlate-zh`` 也不计 ``translated``。"""
    src = _pin_book("<p>Para one words.</p><p>Para two words.</p>", tmp_path, "e")
    dst = tmp_path / "e.out.epub"
    rep = translate_epub(src, dst, _EmptyTranslator())
    out = BeautifulSoup(zipfile.ZipFile(dst).read("OEBPS/ch1.xhtml"), "html.parser")
    zh_nodes = out.select(".texlate-zh")
    assert rep.translated == 0
    assert not zh_nodes

    doc = Document()
    doc.add_paragraph("Docx para words here.")
    dsrc = tmp_path / "e.docx"
    doc.save(str(dsrc))
    ddst = tmp_path / "e.out.docx"
    drep = translate_docx(dsrc, ddst, _EmptyTranslator())
    dparas = list(Document(str(ddst)).element.body.iter(qn("w:p")))
    assert drep.translated == 0
    assert len(dparas) == 1


def test_pre_br_word_boundary(tmp_path: Path) -> None:
    """``<pre>`` 内 ``<br>`` 必须形成词边界（渲染等价换行）。"""
    src = _pin_book("<pre>alpha<br>beta</pre>", tmp_path, "pre")
    book = load_epub(src)
    units = list(iter_units(book, _book_soups(book)))
    texts = [u.text for u in units]
    assert texts
    assert all("alphabeta" not in t for t in texts)
    assert any("alpha" in t and "beta" in t for t in texts)


def test_mimetype_canonical(tmp_path: Path) -> None:
    """非法 mimetype 成员 → 出包必须写规范值。"""
    members = {
        "mimetype": b"TOTALLY WRONG",
        _CONTAINER_PATH: _container_xml("OEBPS/content.opf").encode(),
        "OEBPS/content.opf": _opf(
            [("c0", "ch1.xhtml", "application/xhtml+xml")], ["c0"]
        ).encode(),
        "OEBPS/ch1.xhtml": _xhtml("<p>Mime garbage prose words.</p>"),
    }
    src = _write_epub(tmp_path, "mime", members)
    dst = tmp_path / "mime.out.epub"
    translate_epub(src, dst, MockTranslator())
    assert sniff_format(dst) == "epub"


def test_deep_nesting_recursion(tmp_path: Path) -> None:
    """2000 层行内嵌套：要么翻译成功、要么 ExportError 干净拒绝。"""
    body = f"<p>{'<span>' * _DEEP_NEST}deep{('</span>' * _DEEP_NEST)}</p>"
    src = _pin_book(body, tmp_path, "deep")
    dst = tmp_path / "deep.out.epub"
    try:
        rep = translate_epub(src, dst, MockTranslator())
    except ExportError:
        return  # 干净拒绝也合格
    assert rep.dst == dst


def test_src_zh_class_still_retranslated(tmp_path: Path) -> None:
    """源侧 ``texlate-zh`` 段不豁免枚举——现状钉：照常克隆重译。

    ``_ancestor_skip_reason`` 无类名豁免条款：源文自带 ``texlate-zh`` 的段
    仍进 ``iter_units`` 并被克隆插译——输出 zh 节点 = 源侧原有 + 新插克隆
    （本例 1 + 2 = 3），``translated`` 只计新插（2）。若改判「已译跳过」
    语义此钉须随改。
    """
    src = _pin_book(
        '<p class="texlate-zh">Already translated src.</p>'
        "<p>Fresh para words here.</p>",
        tmp_path,
        "zhsrc",
    )
    book = load_epub(src)
    n_units = len(list(iter_units(book, _book_soups(book))))
    dst = tmp_path / "zhsrc.out.epub"
    rep = translate_epub(src, dst, MockTranslator())
    out = BeautifulSoup(zipfile.ZipFile(dst).read("OEBPS/ch1.xhtml"), "html.parser")
    # 源侧 zh 段照常枚举计数；zh 节点 = 新插克隆 + 源侧自带 1
    assert rep.translated == n_units
    assert len(out.select(".texlate-zh")) == rep.translated + 1
