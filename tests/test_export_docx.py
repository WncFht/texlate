r"""``export.docx`` 双语插译测试——python-docx 现造最小 docx，不碰真文档/真网关。

覆盖：deepcopy ``w:p`` + ``addnext`` 插译、``pPr`` 样式继承、译文 run 的
``w:eastAsia``/颜色戳、echo 不重复插、脚注 part（plain ``Part`` → ``blob``
解析 → ``_blob`` 写回）、``sniff_format``/``export_document`` 分派。
"""

from __future__ import annotations

import re
import zipfile
from typing import TYPE_CHECKING

from docx import Document
from docx.oxml.ns import qn

from texlate.export import export_document, sniff_format
from texlate.export.docx import translate_docx
from texlate.xlat.pipeline import MockTranslator

if TYPE_CHECKING:
    from pathlib import Path

FOOTNOTES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>Footnote body text here.</w:t></w:r></w:p></w:footnote>
</w:footnotes>
"""

_FOOTNOTE_CT = (
    '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
    'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
)
_FOOTNOTE_REL = (
    '<Relationship Id="rIdTexlateFootnotes" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    'footnotes" Target="footnotes.xml"/>'
)


def _make_docx(path: Path, paras: list[tuple[str, str | None]]) -> Path:
    doc = Document()
    for text, style in paras:
        doc.add_paragraph(text, style=style)
    doc.save(str(path))
    return path


def _inject_footnotes(path: Path) -> None:
    """给已存 docx 追加 ``word/footnotes.xml`` part（content-type + rel 全配）。"""
    with zipfile.ZipFile(path) as z:
        members = {i.filename: z.read(i) for i in z.infolist()}
    ct = (
        members["[Content_Types].xml"]
        .decode("utf-8")
        .replace("</Types>", f"{_FOOTNOTE_CT}</Types>")
    )
    rels = (
        members["word/_rels/document.xml.rels"]
        .decode("utf-8")
        .replace("</Relationships>", f"{_FOOTNOTE_REL}</Relationships>")
    )
    members["[Content_Types].xml"] = ct.encode("utf-8")
    members["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    members["word/footnotes.xml"] = FOOTNOTES_XML.encode("utf-8")
    with zipfile.ZipFile(path, "w") as z:
        for name, blob in members.items():
            z.writestr(name, blob, compress_type=zipfile.ZIP_DEFLATED)


class _EchoTranslator:
    """原样回显（剥掉批行 ``[n]`` 前缀）——译文 == 原文走 unchanged 分支。"""

    async def translate(self, **_kw: object) -> str:
        user = str(_kw["user"])
        return "\n".join(re.sub(r"^\[\d+\]\s?", "", ln) for ln in user.split("\n"))


def test_translate_docx_bilingual(tmp_path: Path) -> None:
    src = _make_docx(
        tmp_path / "in.docx",
        [
            ("First paragraph of the document.", None),
            ("List item text here.", "List Number"),
            ("Third and final paragraph.", None),
        ],
    )
    dst = tmp_path / "out.docx"
    report = translate_docx(src, dst, MockTranslator())
    assert report.format == "docx"
    assert report.translated == report.units == 3  # noqa: PLR2004 -- 三段正文

    doc = Document(str(dst))
    paras = doc.paragraphs
    assert len(paras) == 6  # noqa: PLR2004 -- 每段后跟译文克隆
    assert paras[0].text == "First paragraph of the document."
    assert "这是译文" in paras[1].text
    assert paras[2].text == "List item text here."
    assert "这是译文" in paras[3].text

    # 译文 run：eastAsia 字体 + 区分色；列表段 pPr 样式继承
    zh_run = paras[1].runs[0]
    rpr = zh_run._r.find(qn("w:rPr"))  # noqa: SLF001 -- oxml 断言面
    assert rpr is not None
    rfonts = rpr.find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:eastAsia")) == "SimSun"
    color = rpr.find(qn("w:color"))
    assert color is not None
    assert color.get(qn("w:val")) == "555555"

    zh_ppr = paras[3]._p.find(qn("w:pPr"))  # noqa: SLF001 -- oxml 断言面
    assert zh_ppr is not None
    pstyle = zh_ppr.find(qn("w:pStyle"))
    assert pstyle is not None
    assert pstyle.get(qn("w:val")) == "ListNumber"


def test_footnotes_part_translated(tmp_path: Path) -> None:
    src = _make_docx(tmp_path / "in.docx", [("Body paragraph with a note.", None)])
    _inject_footnotes(src)
    dst = tmp_path / "out.docx"
    report = translate_docx(src, dst, MockTranslator())
    assert report.units == 2  # noqa: PLR2004 -- 正文段 + 脚注段
    assert report.documents == 2  # noqa: PLR2004 -- document.xml + footnotes.xml

    with zipfile.ZipFile(dst) as z:
        footnotes = z.read("word/footnotes.xml").decode("utf-8")
    assert "Footnote body text here." in footnotes
    assert "这是译文" in footnotes


def test_echo_translation_not_duplicated(tmp_path: Path) -> None:
    src = _make_docx(tmp_path / "in.docx", [("Only paragraph stays same.", None)])
    dst = tmp_path / "out.docx"
    report = translate_docx(src, dst, _EchoTranslator())
    assert report.unchanged == report.units
    doc = Document(str(dst))
    assert len(doc.paragraphs) == 1  # echo 不插重


def test_sniff_and_dispatch(tmp_path: Path) -> None:
    src = _make_docx(tmp_path / "in.docx", [("Dispatch me please.", None)])
    assert sniff_format(src) == "docx"
    report = export_document(src, None, MockTranslator())
    dst = tmp_path / "in_bilingual.docx"
    assert report.dst == dst
    assert dst.exists()
    doc = Document(str(dst))
    assert len(doc.paragraphs) == 2  # noqa: PLR2004 -- 原段+译文段
