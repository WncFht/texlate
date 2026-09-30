r"""``export.docx`` 隐藏 run 的 ``w:val`` 关态回归——CT_OnOff/ST_OnOff 语义。

``w:vanish``/``w:specVanish`` 是 OnOff 型属性：``w:val`` 缺省 = 开，
``0``/``false``/``off`` 是显式关态——``<w:vanish w:val="0"/>`` 的 run 是
可见的（python-docx 的 ``run.font.hidden = False`` 即序列化成此形），
其文本必须照常送模型；只按"元素在场"判隐藏会把可见文字静默丢掉。
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from docx import Document
from docx.oxml.parser import parse_xml

from texlate.export.docx import iter_units

if TYPE_CHECKING:
    from pathlib import Path

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _doc_with_flagged_run(path: Path, flag_xml: str) -> Path:
    """一段 ``visible prefix `` + 带 ``flag_xml`` rPr 的 run + `` suffix``。"""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("visible prefix ")
    flagged = p.add_run("flagged run text")
    rpr = flagged._r.get_or_add_rPr()  # noqa: SLF001 -- oxml 构造面
    rpr.append(parse_xml(flag_xml))
    p.add_run(" suffix")
    doc.save(str(path))
    return path


@pytest.mark.parametrize("val", ["0", "false", "off"])
def test_hidden_prop_explicit_off_is_visible(tmp_path: Path, val: str) -> None:
    """``<w:vanish w:val="{0,false,off}"/>`` 是显式关态——run 文本照常送模型。"""
    src = _doc_with_flagged_run(
        tmp_path / "in.docx",
        f'<w:vanish xmlns:w="{W_NS}" w:val="{val}"/>',
    )
    doc = Document(str(src))
    texts = [u.text for u, _part, _root in iter_units(doc)]
    assert texts == ["visible prefix flagged run text suffix"]


@pytest.mark.parametrize("val", [None, "1", "true", "on"])
def test_hidden_prop_on_states_stay_hidden(tmp_path: Path, val: str | None) -> None:
    """``w:val`` 缺席（缺省开）与 ``1``/``true``/``on`` 开态——run 仍不送模型。"""
    attr = f' w:val="{val}"' if val is not None else ""
    src = _doc_with_flagged_run(
        tmp_path / "in.docx",
        f'<w:vanish xmlns:w="{W_NS}"{attr}/>',
    )
    doc = Document(str(src))
    texts = [u.text for u, _part, _root in iter_units(doc)]
    assert texts == ["visible prefix suffix"]


def test_specvanish_off_state_is_visible(tmp_path: Path) -> None:
    """``w:specVanish`` 与 ``w:vanish`` 同族——显式关态 run 同样可见。"""
    src = _doc_with_flagged_run(
        tmp_path / "in.docx",
        f'<w:specVanish xmlns:w="{W_NS}" w:val="0"/>',
    )
    doc = Document(str(src))
    texts = [u.text for u, _part, _root in iter_units(doc)]
    assert texts == ["visible prefix flagged run text suffix"]
