"""``texlate export`` 动词：EPUB/DOCX 双语插译 + ``ExportError`` → exit 1 映射。

CLI 层契约：zip 内容嗅探分派（不看后缀）、``--mock``/无 key 默认
MockTranslator 零触网、缺省输出 ``{stem}_bilingual{ext}``、成功报告行落
stderr、``ExportError`` 各族（unsupported/DRM/fixed-layout/malformed）一律
``export: …`` + exit 1 且不产 dst。插译内部语义（克隆/marker/续跑/state）
由 test_export_epub.py / test_export_docx.py 钉，本文件不重复。
"""

from __future__ import annotations

import zipfile
from typing import TYPE_CHECKING

from _exportkit import _epub, _write_epub, _zip_only
from docx import Document
from typer.testing import CliRunner

from texlate.cli import app
from texlate.xlat.pipeline import MOCK_ZH

if TYPE_CHECKING:
    from pathlib import Path

    import pytest

_RUNNER = CliRunner()


def _make_docx(path: Path, paras: list[str]) -> Path:
    doc = Document()
    for text in paras:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


_CH1 = {"ch1.xhtml": "<p>Hello world this is a paragraph for export.</p>"}


class TestExport:
    def test_epub_mock_default_out(self, tmp_path: Path) -> None:
        """--mock → 缺省 ``{stem}_bilingual.epub`` + stderr 报告行 + 无 state 残留。"""
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False))
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 0, result.output
        dst = tmp_path / "book_bilingual.epub"
        assert dst.is_file()
        assert str(dst) in result.stderr
        assert "插译" in result.stderr
        with zipfile.ZipFile(dst) as z:
            ch1 = z.read("OEBPS/ch1.xhtml").decode("utf-8")
        assert "Hello world this is a paragraph" in ch1
        assert MOCK_ZH in ch1
        assert not list(tmp_path.glob("*.state"))  # 断点目录成功即清理

    def test_docx_mock_explicit_out(self, tmp_path: Path) -> None:
        """--mock --out → docx 落指定路径，每原段后跟译文段。"""
        src = _make_docx(
            tmp_path / "in.docx",
            ["First paragraph of the doc.", "Second paragraph follows."],
        )
        dst = tmp_path / "out.docx"
        result = _RUNNER.invoke(app, ["export", str(src), "--mock", "--out", str(dst)])

        assert result.exit_code == 0, result.output
        assert str(dst) in result.stderr
        doc = Document(str(dst))
        paras = doc.paragraphs
        assert len(paras) == 4  # noqa: PLR2004 -- 2 原段 + 2 译文克隆
        assert paras[0].text == "First paragraph of the doc."
        assert MOCK_ZH in paras[1].text

    def test_no_flag_no_key_uses_mock(
        self,
        tmp_path: Path,
        clean_env: pytest.MonkeyPatch,  # noqa: ARG002 -- fixture 副作用（env 清洗）
    ) -> None:
        """无 --mock 且无 TEXLATE_API_KEY → 缺省 MockTranslator（零触网仍成跑），
        隐式回落打 stderr 提示——占位译文不当真译文。"""
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False))
        result = _RUNNER.invoke(app, ["export", str(src)])

        assert result.exit_code == 0, result.output
        assert "MockTranslator" in result.stderr
        assert (tmp_path / "book_bilingual.epub").is_file()

    def test_explicit_mock_no_hint(
        self,
        tmp_path: Path,
        clean_env: pytest.MonkeyPatch,  # noqa: ARG002 -- fixture 副作用（env 清洗）
    ) -> None:
        """显式 --mock 是用户明知 → 不打隐式回落提示。"""
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False))
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 0, result.output
        assert "MockTranslator" not in result.stderr

    def test_gateway_forced_no_key_exit_2(
        self,
        tmp_path: Path,
        clean_env: pytest.MonkeyPatch,
    ) -> None:
        """``TEXLATE_TRANSLATOR=gateway`` 无 key → exit 2 显式拒（必败不静默）。"""
        clean_env.setenv("TEXLATE_TRANSLATOR", "gateway")
        clean_env.delenv("TEXLATE_API_KEY", raising=False)
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False))
        result = _RUNNER.invoke(app, ["export", str(src)])

        assert result.exit_code == 2  # noqa: PLR2004 -- 用法错（配置自相矛盾）
        assert "TEXLATE_API_KEY" in result.stderr
        assert not (tmp_path / "book_bilingual.epub").exists()

    def test_suffix_ignored_content_sniffed(self, tmp_path: Path) -> None:
        """epub 字节命名 ``.bin`` → 内容嗅探照走（help 承诺「不看后缀」）。"""
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False), name="odd.bin")
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 0, result.output
        assert (tmp_path / "odd_bilingual.bin").is_file()

    def test_missing_file_exit_1(self, tmp_path: Path) -> None:
        """不存在的路径 → 嗅探 None → UnsupportedFormatError → exit 1。"""
        result = _RUNNER.invoke(app, ["export", str(tmp_path / "ghost.epub"), "--mock"])

        assert result.exit_code == 1
        assert "export:" in result.stderr
        assert "无法识别" in result.stderr

    def test_unsupported_zip_exit_1(self, tmp_path: Path) -> None:
        """是 zip 但既非 epub 也非 docx → 同一 unsupported 终态。"""
        src = tmp_path / "x.epub"
        src.write_bytes(_zip_only({"readme.txt": "hi"}))
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 1
        assert "无法识别" in result.stderr

    def test_drm_refused_exit_1(self, tmp_path: Path) -> None:
        """META-INF/rights.xml 声明 → DrmError → exit 1 且不产 dst。"""
        src = _write_epub(
            tmp_path,
            _epub(_CH1, ncx=False, extra={"META-INF/rights.xml": "<rights/>"}),
        )
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 1
        assert "protected by DRM" in result.stderr
        assert not (tmp_path / "book_bilingual.epub").exists()

    def test_fixed_layout_exit_1(self, tmp_path: Path) -> None:
        """``rendition:layout=pre-paginated`` → FixedLayoutError → exit 1。"""
        src = _write_epub(tmp_path, _epub(_CH1, ncx=False, fixed=True))
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 1
        assert "fixed-layout" in result.stderr

    def test_malformed_epub_exit_1(self, tmp_path: Path) -> None:
        """mimetype 声明 epub 但缺 container.xml → MalformedEpubError → exit 1。"""
        src = tmp_path / "broken.epub"
        src.write_bytes(_zip_only({"mimetype": "application/epub+zip"}))
        result = _RUNNER.invoke(app, ["export", str(src), "--mock"])

        assert result.exit_code == 1
        assert "container.xml" in result.stderr
